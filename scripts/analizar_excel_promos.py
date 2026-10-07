#!/usr/bin/env python3
"""Analiza el Excel de promociones de Mercado Libre: cuantas publicaciones
ofrecen con aporte propio por debajo de un umbral.

Solo lectura. Opcionalmente exporta a CSV las que cumplen.

Uso:
    venv/bin/python3 scripts/analizar_excel_promos.py <archivo.xlsx>
    venv/bin/python3 scripts/analizar_excel_promos.py <archivo.xlsx> --max 4
    venv/bin/python3 scripts/analizar_excel_promos.py <archivo.xlsx> --max 4 --csv /tmp/ok.csv
"""
import sys
import csv
from collections import Counter, defaultdict

from openpyxl import load_workbook

HOJA = 'Promociones'


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace('%', '').replace('$', '').replace('.', '').replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return None


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    ruta = sys.argv[1]
    tope = float(sys.argv[sys.argv.index('--max') + 1]) if '--max' in sys.argv else 4.0
    csv_out = sys.argv[sys.argv.index('--csv') + 1] if '--csv' in sys.argv else None

    wb = load_workbook(ruta, read_only=True, data_only=True)
    ws = wb[HOJA]
    filas = list(ws.iter_rows(values_only=True))
    wb.close()

    cab = [str(c).strip() if c else '' for c in filas[0]]
    ix = {n: i for i, n in enumerate(cab)}
    req = ['ITEM_ID', 'SELLER_PERCENTAGE', 'PROMO_NAME']
    if any(r not in ix for r in req):
        print(f'Faltan columnas. Encabezados: {cab}')
        return 1

    def g(f, col):
        i = ix.get(col)
        return f[i] if i is not None and i < len(f) else None

    datos = []
    for f in filas[1:]:
        item = g(f, 'ITEM_ID')
        if not item or not str(item).strip().upper().startswith('MLA'):
            continue            # saltea las filas de encabezado/ayuda
        sp = num(g(f, 'SELLER_PERCENTAGE'))
        if sp is None:
            continue
        datos.append({
            'item': str(item).strip(),
            'sku': str(g(f, 'SKU') or '').strip(),
            'titulo': str(g(f, 'TITLE') or '').strip(),
            'promo': str(g(f, 'PROMO_NAME') or '').strip(),
            'seller_pct': sp,
            'meli_pct': num(g(f, 'MELI_PERCENTAGE')) or 0.0,
            'desc_total': num(g(f, 'DISCOUNT_PERCENTAGE')) or 0.0,
            'precio_orig': num(g(f, 'ORIGINAL_PRICE')) or 0,
            'precio_final': num(g(f, 'FINAL_PRICE')) or 0,
            'recibis': num(g(f, 'NEW_RECEIVES')) or 0,
            'estado': str(g(f, 'STATUS') or '').strip(),
        })

    print(f'ARCHIVO: {ruta.split("/")[-1]}')
    print(f'filas con publicacion: {len(datos)}\n')

    print('POR CAMPANIA')
    print('-' * 74)
    por_promo = defaultdict(list)
    for d in datos:
        por_promo[d['promo']].append(d)
    for promo, items in sorted(por_promo.items(), key=lambda kv: -len(kv[1])):
        ok = [i for i in items if i['seller_pct'] < tope]
        print(f'  {promo[:40]:<40} {len(items):>4} ofrecidas   {len(ok):>4} con aporte < {tope}%')

    print('\nPOR TRAMO DE APORTE PROPIO')
    print('-' * 74)
    tramos = Counter()
    for d in datos:
        p = d['seller_pct']
        if p == 0:
            tramos['0% (lo paga todo ML)'] += 1
        elif p < 2:
            tramos['0,1 - 2%'] += 1
        elif p < 3:
            tramos['2 - 3%'] += 1
        elif p < 4:
            tramos['3 - 4%'] += 1
        elif p < 6:
            tramos['4 - 6%'] += 1
        elif p < 10:
            tramos['6 - 10%'] += 1
        else:
            tramos['10% o mas'] += 1
    orden = ['0% (lo paga todo ML)', '0,1 - 2%', '2 - 3%', '3 - 4%',
             '4 - 6%', '6 - 10%', '10% o mas']
    for t in orden:
        if tramos.get(t):
            marca = '  <= te sirven' if t in orden[:4] else ''
            print(f'  {t:<24} {tramos[t]:>4}{marca}')

    # --rango min,max: analiza un tramo puntual (ej. "4,6" = aporte propio 4% a 6%)
    if '--rango' in sys.argv:
        lo, hi = [float(x) for x in sys.argv[sys.argv.index('--rango') + 1].split(',')]
        seg = sorted([d for d in datos if lo <= d['seller_pct'] < hi],
                     key=lambda d: (d['meli_pct'] - d['seller_pct']))
        print(f'\n{"="*74}')
        print(f'TRAMO: aporte propio entre {lo}% y {hi}%   ->  {len(seg)} publicaciones')
        print(f'{"="*74}')
        if not seg:
            return 0
        ml = [d['meli_pct'] for d in seg]
        pr = [d['seller_pct'] for d in seg]
        tot = [d['desc_total'] for d in seg]
        prom_ml = sum(ml) / len(ml)
        prom_pr = sum(pr) / len(pr)
        print(f'  aporte TUYO : min {min(pr):.2f}%  prom {prom_pr:.2f}%  max {max(pr):.2f}%')
        print(f'  aporte de ML: min {min(ml):.2f}%  prom {prom_ml:.2f}%  max {max(ml):.2f}%')
        print(f'  descuento total: prom {sum(tot)/len(tot):.1f}%')
        print(f'  de cada 100% de descuento, ML pone {prom_ml/(prom_ml+prom_pr)*100:.0f}% '
              f'y vos {prom_pr/(prom_ml+prom_pr)*100:.0f}%')

        cnt = Counter()
        for d in seg:
            m = d['meli_pct']
            if m == 0:
                cnt['ML no pone nada (0%)'] += 1
            elif m < 1:
                cnt['ML pone menos de 1%'] += 1
            elif m < 2:
                cnt['ML pone 1 - 2%'] += 1
            elif m < 4:
                cnt['ML pone 2 - 4%'] += 1
            elif m < 6:
                cnt['ML pone 4 - 6%'] += 1
            else:
                cnt['ML pone 6% o mas'] += 1
        print('\n  CUANTO PONE ML EN ESE TRAMO')
        for k in ['ML no pone nada (0%)', 'ML pone menos de 1%', 'ML pone 1 - 2%',
                  'ML pone 2 - 4%', 'ML pone 4 - 6%', 'ML pone 6% o mas']:
            if cnt.get(k):
                print(f'    {k:<24} {cnt[k]:>4}   ({cnt[k]/len(seg)*100:.0f}%)')

        mejores = sorted(seg, key=lambda d: -(d['meli_pct'] - d['seller_pct']))[:12]
        print('\n  Las 12 donde ML pone MAS que vos:')
        print(f'  {"MLA":<15} {"SKU":<12} {"tuyo":>6} {"ML":>6} {"dif":>6} {"total":>7}  producto')
        for d in mejores:
            print(f'  {d["item"]:<15} {d["sku"][:12]:<12} {d["seller_pct"]:>5.2f}% '
                  f'{d["meli_pct"]:>5.2f}% {d["meli_pct"]-d["seller_pct"]:>+5.2f} '
                  f'{d["desc_total"]:>6.1f}%  {d["titulo"][:38]}')
        if csv_out:
            with open(csv_out, 'w', newline='', encoding='utf-8-sig') as fh:
                w = csv.DictWriter(fh, fieldnames=list(seg[0].keys()), delimiter=';')
                w.writeheader()
                w.writerows(sorted(seg, key=lambda d: -(d['meli_pct'] - d['seller_pct'])))
            print(f'\nCSV con las {len(seg)}: {csv_out}')
        return 0

    ok = sorted([d for d in datos if d['seller_pct'] < tope], key=lambda d: d['seller_pct'])
    estados = Counter(d['estado'] for d in ok)
    print(f'\n{"="*74}')
    print(f'RESULTADO: {len(ok)} de {len(datos)} publicaciones con aporte propio < {tope}%'
          f'   ({len(ok)/len(datos)*100:.0f}%)')
    print(f'{"="*74}')
    print(f'  estado de esas: {dict(estados)}')
    if ok:
        print(f'  aporte propio: min {ok[0]["seller_pct"]:.2f}%  max {ok[-1]["seller_pct"]:.2f}%')
        print(f'  descuento total promedio: {sum(d["desc_total"] for d in ok)/len(ok):.1f}%')
        print(f'  aporte de ML promedio   : {sum(d["meli_pct"] for d in ok)/len(ok):.1f}%')

        print(f'\n  Las 15 mas convenientes (menor aporte tuyo):')
        print(f'  {"MLA":<15} {"SKU":<12} {"tuyo":>6} {"ML":>6} {"total":>7}  producto')
        for d in ok[:15]:
            print(f'  {d["item"]:<15} {d["sku"][:12]:<12} {d["seller_pct"]:>5.2f}% '
                  f'{d["meli_pct"]:>5.2f}% {d["desc_total"]:>6.1f}%  {d["titulo"][:42]}')

    if csv_out and ok:
        with open(csv_out, 'w', newline='', encoding='utf-8-sig') as fh:
            w = csv.DictWriter(fh, fieldnames=list(ok[0].keys()), delimiter=';')
            w.writeheader()
            w.writerows(ok)
        print(f'\nCSV con las {len(ok)}: {csv_out}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
