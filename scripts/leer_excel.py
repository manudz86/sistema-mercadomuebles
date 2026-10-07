#!/usr/bin/env python3
"""Inspecciona un Excel: hojas, encabezados y primeras filas. Solo lectura.

Uso:
    venv/bin/python3 scripts/leer_excel.py <archivo.xlsx>
    venv/bin/python3 scripts/leer_excel.py <archivo.xlsx> --filas 20
"""
import sys

from openpyxl import load_workbook


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    ruta = sys.argv[1]
    nfilas = int(sys.argv[sys.argv.index('--filas') + 1]) if '--filas' in sys.argv else 6

    wb = load_workbook(ruta, read_only=True, data_only=True)
    print(f'hojas: {wb.sheetnames}\n')
    for hoja in wb.sheetnames:
        ws = wb[hoja]
        print(f'── {hoja}  ({ws.max_row} filas x {ws.max_column} columnas)')
        for i, fila in enumerate(ws.iter_rows(values_only=True), 1):
            if i > nfilas:
                break
            vals = ['' if v is None else str(v)[:26] for v in fila]
            print(f'  {i:>3}: ' + ' | '.join(vals))
        print()
    wb.close()
    return 0


if __name__ == '__main__':
    sys.exit(main())
